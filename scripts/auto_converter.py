#!/usr/bin/env python3
"""
MdMax Auto-Converter - Automatic .xls → .xlsx conversion
The COMPETITIVE ADVANTAGE: Convert legacy formats silently without user intervention
"""

import tempfile
from pathlib import Path
from typing import Tuple, Optional


def auto_convert_xls_to_xlsx(file_path: str) -> Tuple[str, bool]:
    """
    Automatically convert .xls to .xlsx format silently

    Returns:
        Tuple of (converted_file_path, was_converted)
        - If already .xlsx: returns original path, False
        - If .xls: converts and returns temp .xlsx path, True

    This is the COMPETITIVE ADVANTAGE:
    - Zero user intervention
    - Automatic format detection and conversion
    - Seamless processing
    """

    file_path_obj = Path(file_path)

    # If already xlsx, return as-is
    if file_path_obj.suffix.lower() == ".xlsx":
        return file_path, False

    # If .xls, attempt automatic conversion
    if file_path_obj.suffix.lower() == ".xls":
        try:
            # Method 1: Try pandas (handles both formats elegantly)
            import pandas as pd

            try:
                # Read .xls file
                df_dict = pd.read_excel(file_path, sheet_name=None)

                # Create temporary .xlsx file
                temp_xlsx = tempfile.NamedTemporaryFile(
                    suffix=".xlsx",
                    delete=False
                )
                temp_xlsx_path = temp_xlsx.name
                temp_xlsx.close()

                # Write to .xlsx using pandas
                with pd.ExcelWriter(temp_xlsx_path, engine='openpyxl') as writer:
                    for sheet_name, df in df_dict.items():
                        df.to_excel(writer, sheet_name=sheet_name, index=False)

                return temp_xlsx_path, True

            except Exception as e:
                print(f"[INFO] Pandas conversion failed: {e}")
                # Fallback to method 2
                pass

        except ImportError:
            print(f"[INFO] Pandas not installed, trying alternative method")
            pass

        # Method 2: Try xlrd + openpyxl combination
        try:
            import xlrd
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter

            # Read .xls file
            workbook = xlrd.open_workbook(file_path)

            # Create new xlsx workbook
            new_wb = Workbook()
            new_wb.remove(new_wb.active)  # Remove default sheet

            # Copy all sheets
            for sheet_name in workbook.sheet_names():
                sheet = workbook.sheet_by_name(sheet_name)
                new_sheet = new_wb.create_sheet(title=sheet_name)

                # Copy all cells
                for row_idx in range(sheet.nrows):
                    for col_idx in range(sheet.ncols):
                        cell_value = sheet.cell_value(row_idx, col_idx)
                        new_sheet.cell(
                            row=row_idx + 1,
                            column=col_idx + 1,
                            value=cell_value
                        )

            # Save to temporary .xlsx
            temp_xlsx = tempfile.NamedTemporaryFile(
                suffix=".xlsx",
                delete=False
            )
            temp_xlsx_path = temp_xlsx.name
            temp_xlsx.close()

            new_wb.save(temp_xlsx_path)

            return temp_xlsx_path, True

        except Exception as e:
            print(f"[INFO] xlrd + openpyxl conversion failed: {e}")
            pass

    # If conversion not possible, return original
    return file_path, False


def auto_detect_and_convert(file_path: str) -> str:
    """
    Automatically detect format and convert if needed.
    Returns path to file ready for processing (either original or converted).

    This is the MAGIC of MdMax v2.1:
    - User sends .xls
    - MdMax automatically converts to .xlsx
    - User receives perfect Markdown
    - Zero manual steps
    """

    converted_path, was_converted = auto_convert_xls_to_xlsx(file_path)

    if was_converted:
        print(f"[AUTO-CONVERT] 🔄 Automatically converted .xls → .xlsx")
        print(f"[AUTO-CONVERT] Original: {file_path}")
        print(f"[AUTO-CONVERT] Converted: {converted_path}")

    return converted_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1:
        file_path = sys.argv[1]
        converted = auto_detect_and_convert(file_path)
        print(f"Ready to process: {converted}")
    else:
        print("Usage: python auto_converter.py <file_path>")
