#!/usr/bin/env python3
"""
SmartConverter: Analisa arquivos e escolhe melhor formato para conversão.

Detecta automaticamente se um arquivo é "denso" (muitos dados) e recomenda:
- Markdown para arquivos esparsos/pequenos
- JSON compactado para arquivos densos/grandes

Resolve problema onde arquivos > 500KB ficam MAIORES em Markdown.
"""

import json
import os
from pathlib import Path
from openpyxl import load_workbook


class SmartConverter:
    """Analisador de densidade e recomendador de formato de conversão."""

    def __init__(self):
        self.density_threshold = 0.70  # 70% de células preenchidas
        self.size_threshold = 500_000  # 500KB

    def analyze_excel(self, file_path):
        """
        Analisa arquivo Excel e retorna métricas.

        Returns:
            dict: {
                'file_size': int (bytes),
                'density': float (0.0-1.0),
                'recommendation': 'JSON' | 'MARKDOWN',
                'reason': str
            }
        """
        file_path = Path(file_path)
        file_size = file_path.stat().st_size

        try:
            wb = load_workbook(file_path, data_only=True)
            ws = wb.active

            # Contar células preenchidas
            total_cells = 0
            filled_cells = 0

            for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                   min_col=1, max_col=ws.max_column):
                for cell in row:
                    total_cells += 1
                    if cell.value is not None:
                        filled_cells += 1

            density = filled_cells / total_cells if total_cells > 0 else 0

            # Recomendação baseada em densidade E tamanho
            is_dense = density > self.density_threshold
            is_large = file_size > self.size_threshold

            if is_dense or is_large:
                recommendation = 'JSON'
                reason = []
                if is_dense:
                    reason.append(f'densidade alta ({density*100:.0f}%)')
                if is_large:
                    reason.append(f'tamanho grande ({file_size/1024/1024:.1f}MB)')
                reason_str = ' + '.join(reason)
            else:
                recommendation = 'MARKDOWN'
                reason_str = f'dados esparsos ({density*100:.0f}%) e pequeno ({file_size/1024:.0f}KB)'

            return {
                'file_size': file_size,
                'file_size_mb': file_size / 1024 / 1024,
                'density': density,
                'filled_cells': filled_cells,
                'total_cells': total_cells,
                'recommendation': recommendation,
                'reason': reason_str
            }

        except Exception as e:
            return {
                'error': str(e),
                'file_size': file_size,
                'recommendation': 'MARKDOWN',  # fallback seguro
                'reason': 'erro na análise, usando Markdown como fallback'
            }

    def excel_to_json(self, file_path, output_path):
        """
        Converte Excel para JSON compactado.

        Usa separadores mínimos para máxima compressão.
        """
        file_path = Path(file_path)
        output_path = Path(output_path)

        wb = load_workbook(file_path, data_only=True)
        ws = wb.active

        # Extrair headers (primeira linha)
        headers = []
        for cell in ws[1]:
            headers.append(cell.value or "")

        # Extrair dados
        rows = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            row_data = {}
            for col_idx, value in enumerate(row):
                if col_idx < len(headers):
                    row_data[headers[col_idx]] = value
            rows.append(row_data)

        # Criar estrutura JSON
        data = {
            'source': file_path.name,
            'format': 'converted_excel',
            'columns': headers,
            'row_count': len(rows),
            'data': rows
        }

        # Salvar com máxima compressão
        with open(output_path, 'w', encoding='utf-8') as f:
            # separators=(',', ':') remove espaços em branco desnecessários
            json.dump(data, f, separators=(',', ':'), ensure_ascii=False)

        return {
            'output_file': str(output_path),
            'rows_converted': len(rows),
            'output_size': output_path.stat().st_size
        }


if __name__ == '__main__':
    import sys

    if len(sys.argv) < 2:
        print("Uso: python smart_converter.py <arquivo.xlsx>")
        sys.exit(1)

    converter = SmartConverter()
    file_path = sys.argv[1]

    # Análise
    result = converter.analyze_excel(file_path)

    print("\n=== ANÁLISE DE CONVERSÃO ===\n")
    print(f"Arquivo: {file_path}")
    print(f"Tamanho: {result.get('file_size_mb', 0):.1f}MB")
    print(f"Densidade: {result.get('density', 0)*100:.0f}%")
    print(f"Recomendação: {result.get('recommendation')}")
    print(f"Motivo: {result.get('reason')}")
    print()
